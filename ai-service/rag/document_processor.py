import os
from typing import List, Dict, Any
from loguru import logger

class DocumentProcessor:
    """
    Extracts text and chunks documents for CSE RAG Pipeline.
    Supports PDF, DOCX, PPTX, and TXT files.
    """
    def __init__(self, chunk_size: int = 600, chunk_overlap: int = 80):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def extract_text(self, file_path: str) -> str:
        """
        Extract raw text content from the file based on its extension.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        try:
            if ext == '.pdf':
                return self._extract_pdf(file_path)
            elif ext in ['.docx', '.doc']:
                return self._extract_docx(file_path)
            elif ext in ['.pptx', '.ppt']:
                return self._extract_pptx(file_path)
            else:
                # Text or fallback
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
        except Exception as e:
            logger.error(f"[DocumentProcessor] Error extracting text from {file_path}: {e}")
            return f"Document content from {os.path.basename(file_path)}."

    def _extract_pdf(self, file_path: str) -> str:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"--- Page {i+1} ---\n{page_text}")
        return "\n\n".join(text_parts)

    def _extract_docx(self, file_path: str) -> str:
        import docx
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)

    def _extract_pptx(self, file_path: str) -> str:
        from pptx import Presentation
        prs = Presentation(file_path)
        slides_text = []
        for i, slide in enumerate(prs.slides):
            slide_content = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())
            if slide_content:
                slides_text.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_content))
        return "\n\n".join(slides_text)

    def chunk_text(self, text: str, metadata: Dict[str, Any] = None) -> List[Dict[str, Any]]:
        """
        Split text into overlapping chunks and attach document metadata.
        """
        if not text:
            return []

        chunks = []
        start = 0
        text_length = len(text)
        chunk_idx = 0

        while start < text_length:
            end = min(start + self.chunk_size, text_length)
            
            # If not at the end of the text, try to break at a newline or space
            if end < text_length:
                break_point = text.rfind('\n', start, end)
                if break_point == -1 or break_point <= start:
                    break_point = text.rfind(' ', start, end)
                if break_point > start:
                    end = break_point

            chunk_content = text[start:end].strip()
            if chunk_content:
                chunks.append({
                    "chunk_id": chunk_idx,
                    "content": chunk_content,
                    "metadata": {
                        **(metadata or {}),
                        "chunk_index": chunk_idx,
                        "char_start": start,
                        "char_end": end
                    }
                })
                chunk_idx += 1

            start = end - self.chunk_overlap
            if start >= end:
                start = end

        return chunks

document_processor = DocumentProcessor()
